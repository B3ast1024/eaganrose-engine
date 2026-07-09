import os
import streamlit as st
import pdfplumber
import pandas as pd
import re
from datetime import datetime
import io
from dateutil.relativedelta import relativedelta
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
import plotly.express as px

# =====================================================================
# 1. AUTO-THEMING ENGINE & SESSION STATE
# =====================================================================
os.makedirs(".streamlit", exist_ok=True)
theme_config = """
[theme]
primaryColor = "#8E2D28"
backgroundColor = "#FFFFFF"
secondaryBackgroundColor = "#F0F2F6"
textColor = "#1E1E1E"
font = "sans serif"
"""
with open(".streamlit/config.toml", "w") as f:
    f.write(theme_config)

st.set_page_config(page_title="Eaganrose Platform", layout="wide")

# Initialize memory for Login, Audit Logging, and the Advanced CRM
if "logged_in" not in st.session_state:
    st.session_state["logged_in"] = False
if "username" not in st.session_state:
    st.session_state["username"] = ""
if "audit_log" not in st.session_state:
    st.session_state["audit_log"] = []
    
# UPGRADED DATA STRUCTURE: Now supports default spoils and SKU-level overrides
if "vendor_registry" not in st.session_state:
    st.session_state["vendor_registry"] = {
        "0010006727": {
            "name": "DAESANG",
            "default_spoils": 0.0075, # 0.75% default for Daesang
            "skus": {
                "1570571": {"comm": 0.03, "spoils": None}, 
                "1793017": {"comm": 0.03, "spoils": None}, 
                "1908395": {"comm": 0.05, "spoils": None}, 
                "1990882": {"comm": 0.05, "spoils": None}, 
                "2059134": {"comm": 0.05, "spoils": None}
            }
        }
    }

# =====================================================================
# 2. AUTHENTICATION MODULE
# =====================================================================
USER_DATABASE = {
    "colton.rose": "admin123",
    "account.manager": "eaganrose"
}

def login_screen():
    st.markdown("<h1 style='text-align: center;'>Eaganrose Secure Portal</h1>", unsafe_allow_html=True)
    st.markdown("---")
    
    col1, col2, col3 = st.columns([1, 1, 1])
    with col2:
        st.markdown("### 🔐 Authorized Access Only")
        username_input = st.text_input("Username")
        password_input = st.text_input("Password", type="password")
        
        if st.button("Sign In", use_container_width=True):
            if username_input in USER_DATABASE and USER_DATABASE[username_input] == password_input:
                st.session_state["logged_in"] = True
                st.session_state["username"] = username_input
                st.rerun()
            else:
                st.error("Invalid credentials. Please try again.")

if not st.session_state["logged_in"]:
    login_screen()
    st.stop()

# =====================================================================
# 3. GLOBAL SIDEBAR (AUDIT LOG & LOGOUT)
# =====================================================================
with st.sidebar:
    st.success(f"👤 Logged in as: **{st.session_state['username']}**")
    st.markdown("---")
    
    st.markdown("### 📜 Session Audit Log")
    if not st.session_state["audit_log"]:
        st.info("No actions taken this session.")
    else:
        for entry in reversed(st.session_state["audit_log"][-5:]):
            st.caption(f"**{entry['timestamp']}** | 👤 {entry['user']}")
            st.write(f"Processed {entry['files_count']} POs.")
            st.markdown("---")
            
    if st.button("Log Out", use_container_width=True):
        st.session_state["logged_in"] = False
        st.session_state["username"] = ""
        st.rerun()

# =====================================================================
# 4. CORE EXTRACTION ENGINE (UPGRADED FOR DYNAMIC ALLOWANCES & REVISIONS)
# =====================================================================
class InvalidCostcoPOError(Exception):
    pass

def extract_costco_pdf_data(pdf_file_obj):
    text = ""
    try:
        with pdfplumber.open(pdf_file_obj) as pdf:
            for page in pdf.pages:
                extracted = page.extract_text()
                if extracted:
                    text += extracted + "\n"
    except Exception as e:
        raise InvalidCostcoPOError(f"File is corrupted or not a readable PDF. ({str(e)})")
        
    order_match = re.search(r"Order\s*#[:\s\n]*(\d+)", text, re.IGNORECASE)
    if not order_match:
        raise InvalidCostcoPOError("Missing 'Order #'. This does not appear to be a valid Costco PO.")
        
    header_data = {}
    po_number = order_match.group(1)
    header_data['PO #'] = po_number
    header_data['Purchase Order #'] = po_number 
    
    detected_vendor_name = "UNKNOWN_VENDOR"
    detected_vendor_info = None
    
    # Dynamic Registry Check
    for v_id, v_info in st.session_state["vendor_registry"].items():
        if v_id in text:
            detected_vendor_name = v_info["name"]
            detected_vendor_info = v_info
            break
            
    header_data['Vendor'] = detected_vendor_name
    
    if len(po_number) >= 11:
        header_data['Extracted PO Date'] = po_number[-7:-3] 
        header_data['PO Sequence'] = po_number[-3:]
        header_data['Extracted Depot'] = po_number[:-7].lstrip('0') 
    else:
        header_data['Extracted PO Date'], header_data['PO Sequence'], header_data['Extracted Depot'] = "", "", ""

    # Region / Revision Detection
    dept_match = re.search(r"Department\s*#[:\s\n]*([^\n]+)", text, re.IGNORECASE)
    if dept_match:
        parts = [p.strip() for p in dept_match.group(1).split('/')]
        raw_reg = parts[1] if len(parts) >= 2 else parts[0]
        reg_match = re.search(r"([A-Z]{2})", raw_reg)
        
        # Explicit check for "0" department signifying a revision
        if "0" in raw_reg and not reg_match:
            header_data['Region'] = "0"
        else:
            header_data['Region'] = reg_match.group(1) if reg_match else raw_reg[:2]
    else:
        header_data['Region'] = ""
    
    dates = re.findall(r"(\d{2}/\d{2}/\d{4})", text)
    if len(dates) >= 2:
        po_dt = datetime.strptime(dates[0], "%m/%d/%Y")
        del_dt = datetime.strptime(dates[1], "%m/%d/%Y")
        header_data['PO Date'] = po_dt.strftime("%Y-%m-%d")
        header_data[' PO Del Date'] = del_dt.strftime("%Y-%m-%d")
        header_data['Cancel Date'] = del_dt.strftime("%Y-%m-%d") 
        est_pay = (del_dt + relativedelta(months=1)).replace(day=1)
        header_data["Est'd  Month   Pay'd"] = est_pay.strftime("%Y-%m-%d")
    else:
        header_data['PO Date'], header_data[' PO Del Date'], header_data['Cancel Date'], header_data["Est'd  Month   Pay'd"] = "", "", "", ""
    
    items_data = []
    item_matches = list(re.finditer(r"^(\d{1,3})\s+(\d{7})\s+.*?([A-Za-z].*?)\s+([\d\,\.]+)\s+(\d+)\s+CA", text, re.MULTILINE))
    
    if not item_matches:
        raise InvalidCostcoPOError("Could not isolate line items. The PDF formatting may be unexpected.")
        
    for i, match in enumerate(item_matches):
        item_dict = header_data.copy()
        sku = match.group(2)
        desc = match.group(3).strip().title()
        unit_cost = float(match.group(4).replace(',', ''))
        qty = int(match.group(5))
        
        item_dict['Costco Item #'] = sku
        item_dict['Item Description'] = desc
        item_dict['Unit Cost'] = unit_cost
        item_dict['Qty'] = qty
        
        # --- DYNAMIC RATE APPLICATION ---
        comm_rate = 0.0
        spoils_rate = 0.0075 # Absolute fallback
        
        if detected_vendor_info:
            spoils_rate = detected_vendor_info.get("default_spoils", 0.0075)
            sku_data = detected_vendor_info["skus"].get(sku, {})
            
            if isinstance(sku_data, dict):
                comm_rate = sku_data.get("comm", 0.0)
                if sku_data.get("spoils") is not None:
                    spoils_rate = sku_data["spoils"] # Override with SKU specific spoil rate
            else:
                comm_rate = sku_data # Legacy fallback just in case
                
        item_dict['Com   %   Rate'] = comm_rate
        item_dict['Spoils'] = spoils_rate
        # ---------------------------------
        
        chunk_start = match.end()
        chunk_end = item_matches[i+1].start() if i + 1 < len(item_matches) else len(text)
        item_chunk = text[chunk_start:chunk_end]
        
        demo_total = 0.0
        allowance_lines = re.findall(r"^.*(?:ALLOWANCE|DFI\$).*$", item_chunk, re.IGNORECASE | re.MULTILINE)
        
        for line in allowance_lines:
            if re.search(r"SPOIL", line, re.IGNORECASE):
                continue
            amounts = re.findall(r"[\d\,]+\.\d{2}", line)
            if amounts:
                first_amount = float(amounts[0].replace(',', ''))
                last_amount = float(amounts[-1].replace(',', ''))
                if first_amount < 2.0:
                    continue
                demo_total += last_amount
                    
        item_dict['Demo Accrual Deduction'] = demo_total
        
        freight_match = re.search(r"FREIGHT.*?(\d+\.\d{2})", item_chunk, re.IGNORECASE)
        freight_allowance = float(freight_match.group(1)) * qty if freight_match else 0.0
        item_dict['Freight Allowance'] = freight_allowance
        
        gross_amt = unit_cost * qty
        spoils_calc = gross_amt * spoils_rate # NOW USES DYNAMIC RATE
        item_dict['Net Inv Amt'] = gross_amt - demo_total - spoils_calc - freight_allowance
        
        items_data.append(item_dict)
        
    return items_data

def process_and_merge(df):
    bd_mask = df['Region'].astype(str).str.contains('BD', case=False, na=False)
    bd_data = df[bd_mask].copy()
    d12_data = df[~bd_mask].copy()
    
    if not bd_data.empty:
        bd_data['Dead Net Cost'] = bd_data['Unit Cost'] - bd_data['Freight Allowance'] - bd_data['Demo Accrual Deduction']
        bd_data['Total Dead Net Amt'] = bd_data['Dead Net Cost'] * bd_data['Qty']
        
    master_df = pd.concat([bd_data, d12_data], ignore_index=True)
    if 'PO Date' in master_df.columns:
        master_df = master_df.sort_values(by='PO Date', ascending=True)
    return master_df

def parse_date(date_str):
    if not date_str: return None
    try: return datetime.strptime(date_str, "%Y-%m-%d").date()
    except: return date_str

def create_excel_buffer(master_df):
    output = io.BytesIO()
    wb = Workbook()
    ws = wb.active
    ws.title = "Master_Ledger" 
    
    arial_font = Font(name='Arial', size=8)
    header_font = Font(name='Arial', size=8, bold=True, color="000000")
    header_fill = PatternFill(start_color="FFC000", end_color="FFC000", fill_type="solid")
    thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")
    right_align = Alignment(horizontal="right", vertical="center")
    
    headers = [
        'PO Date', 'Vendor', 'Region', 'Purchase Order #', '', '', '', 
        'Costco Item #', 'Item Description', 'Cancel Date', ' PO Del Date', 
        'Qty', 'Unit Cost', 'Spoils', 'Freight Allowance', 'Demo Accrual Deduction', 
        'Net Inv Amt', 'Com   %   Rate', 'Total Comision Earned', 
        "Est'd  Month   Pay'd", 'Date Comm  Paid', 'Ck #', 'Total Comision Paid', 
        'Commission Paid Difference from Commission Earned', 'Comments'
    ]
    
    ws.append(headers)
    for col_num, cell in enumerate(ws[1], 1):
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align
        cell.border = thin_border
        
    for idx, row in master_df.iterrows():
        ws.append([
            parse_date(row.get('PO Date', '')), row.get('Vendor', ''), row.get('Region', ''),
            row.get('Extracted Depot', ''), row.get('Extracted PO Date', ''), row.get('PO Sequence', ''), '',
            row.get('Costco Item #', ''), row.get('Item Description', ''), parse_date(row.get('Cancel Date', '')),
            parse_date(row.get(' PO Del Date', '')), row.get('Qty', 0), row.get('Unit Cost', 0.0),
            row.get('Spoils', 0.0), row.get('Freight Allowance', 0.0), row.get('Demo Accrual Deduction', 0.0),
            row.get('Net Inv Amt', 0.0), row.get('Com   %   Rate', 0.0), 0.0,
            parse_date(row.get("Est'd  Month   Pay'd", '')), None, None, None, 0.0, ''
        ])
        
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=25):
        for cell in row:
            cell.font = arial_font
            cell.border = thin_border
            cell.alignment = center_align 
            
            col = cell.column
            if col in [4, 24]: cell.alignment = right_align
            elif col in [6, 9]: cell.alignment = left_align
            
            if col == 5: cell.number_format = '0000' 
            elif col in [13, 23]: cell.number_format = '0.00'
            elif col == 12: cell.number_format = '#,##0' 
            elif col in [14, 18]: cell.number_format = '0.00%'
            elif col in [15, 16]: cell.number_format = '_("$"* #,##0.00_);_("$"* (#,##0.00);_("$"* "-"??_);_(@_)'
            elif col in [17, 19, 24]: cell.number_format = '"$"#,##0.00'
            elif col in [1, 10, 11]: cell.number_format = 'dd-mmm-yy'
            elif col == 20: cell.number_format = 'mmm-yy'
                
        row_idx = row[0].row
        ws[f'S{row_idx}'] = f'=Q{row_idx}*R{row_idx}'
        ws[f'X{row_idx}'] = f'=W{row_idx}-S{row_idx}'

    wb.save(output)
    output.seek(0)
    return output


# =====================================================================
# 5. TABBED UI LAYOUT
# =====================================================================
tab_audit, tab_crm = st.tabs(["🔍 Audit Engine", "🤝 Vendor & Allowance Matrix"])

# --- MODULE A: AUDIT ENGINE ---
with tab_audit:
    try:
        st.image("eagan rose logo.png", width=250)
    except Exception:
        pass

    st.title("Eaganrose Reconciliation Engine")
    st.markdown("Automated PDF Ingestion & Variance Auditing for Costco Vendors")
    st.markdown("---")

    uploaded_files = st.file_uploader("Upload Costco POs (PDF)", type="pdf", accept_multiple_files=True)

    if uploaded_files:
        if st.button("Run Audit Engine"):
            with st.spinner(f"Ingesting {len(uploaded_files)} Purchase Orders..."):
                all_data, failed_files = [], []
                for pdf_file in uploaded_files:
                    try:
                        extracted_items = extract_costco_pdf_data(pdf_file)
                        all_data.extend(extracted_items)
                    except Exception as e:
                        failed_files.append({"filename": pdf_file.name, "error": str(e)})
                
                if all_data:
                    master_df = process_and_merge(pd.DataFrame(all_data))
                    
                    unique_pos = list(master_df['Purchase Order #'].unique())
                    st.session_state["audit_log"].append({
                        "timestamp": datetime.now().strftime("%Y-%m-%d %I:%M %p"),
                        "user": st.session_state["username"],
                        "files_count": len(uploaded_files),
                        "po_numbers": ", ".join(unique_pos)
                    })
                    
                    st.success(f"Audit Complete! Processed {len(master_df)} valid records.")
                    
                    # --- REVISION ALERT MODULE ---
                    # Check for any rows where the Region was logged as "0"
                    revisions_df = master_df[master_df['Region'].astype(str).str.strip() == '0']
                    if not revisions_df.empty:
                        rev_po_list = revisions_df['Purchase Order #'].unique().tolist()
                        st.warning(f"⚠️ **Attention: {len(rev_po_list)} Revision(s) Detected!**\nThe following POs were flagged as revisions (Department '0') and will need to be manually updated in your master ledger: **{', '.join(rev_po_list)}**")
                    # -----------------------------
                    
                    st.markdown("### 📊 Operational Summary")
                    col1, col2, col3 = st.columns(3)
                    tot_net = master_df['Net Inv Amt'].sum()
                    tot_comm = (master_df['Net Inv Amt'] * master_df['Com   %   Rate']).sum()
                    tot_qty = int(master_df['Qty'].sum())
                    
                    col1.metric("Total Net Inventory", f"${tot_net:,.2f}")
                    col2.metric("Est. Commission Earned", f"${tot_comm:,.2f}")
                    col3.metric("Total Cases (Qty)", f"{tot_qty:,}")
                    
                    st.markdown("---")
                    chart_col1, chart_col2 = st.columns(2)
                    
                    with chart_col1:
                        st.markdown("**Allowance Breakdown**")
                        # Spoils are now calculated dynamically per row!
                        tot_spoils = (master_df['Unit Cost'] * master_df['Qty'] * master_df['Spoils']).sum()
                        tot_demo = master_df['Demo Accrual Deduction'].sum()
                        tot_freight = master_df['Freight Allowance'].sum()
                        
                        allowance_df = pd.DataFrame({
                            "Allowance Type": ["Spoils", "Demo Accrual", "Freight"],
                            "Amount ($)": [tot_spoils, tot_demo, tot_freight]
                        })
                        
                        fig_donut = px.pie(allowance_df, values='Amount ($)', names='Allowance Type', hole=0.4,
                                           color='Allowance Type',
                                           color_discrete_map={"Spoils": "#8E2D28", "Demo Accrual": "#FFC000", "Freight": "#1E1E1E"})
                        fig_donut.update_layout(margin=dict(t=20, b=20, l=20, r=20))
                        st.plotly_chart(fig_donut, use_container_width=True)
                    
                    with chart_col2:
                        st.markdown("**Volume by SKU (Cases)**")
                        sku_df = master_df.groupby("Item Description")['Qty'].sum().reset_index()
                        sku_df = sku_df.sort_values(by='Qty', ascending=True) 
                        
                        fig_bar = px.bar(sku_df, x='Qty', y='Item Description', orientation='h',
                                         color_discrete_sequence=["#8E2D28"])
                        fig_bar.update_layout(margin=dict(t=20, b=20, l=20, r=20), xaxis_title="Total Cases", yaxis_title="")
                        st.plotly_chart(fig_bar, use_container_width=True)
                    
                    st.markdown("---")
                    st.download_button(
                        label="📥 Download Master Ledger (Excel)",
                        data=create_excel_buffer(master_df),
                        file_name="Audit Software: Delete After Use.xlsx",
                        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                    )
                    
                if failed_files:
                    st.warning(f"⚠️ Skipped {len(failed_files)} file(s).")
                    with st.expander("View Error Log"):
                        for fail in failed_files: st.write(f"- **{fail['filename']}**: {fail['error']}")

# --- MODULE B: VENDOR & ALLOWANCE MATRIX ---
with tab_crm:
    st.title("Vendor & Allowance Matrix")
    st.markdown("Manage vendor IDs, default allowances, and map SKU commission rates dynamically. Changes applied here immediately take effect in the Audit Engine.")
    st.markdown("---")
    
    col_crm_1, col_crm_2 = st.columns(2)
    
    with col_crm_1:
        st.markdown("### ➕ Add/Edit Vendor Profile")
        with st.form("add_vendor_form", clear_on_submit=True):
            new_v_name = st.text_input("Vendor Name (e.g., KRAFT HEINZ)").upper()
            new_v_id = st.text_input("Costco Vendor ID (Exact match required)")
            new_v_spoils = st.number_input("Default Spoils Allowance % (e.g., 0.75)", value=0.75, step=0.05)
            
            submit_vendor = st.form_submit_button("Save Vendor Profile")
            
            if submit_vendor:
                if new_v_id and new_v_name:
                    if new_v_id not in st.session_state["vendor_registry"]:
                        st.session_state["vendor_registry"][new_v_id] = {"name": new_v_name, "default_spoils": new_v_spoils / 100.0, "skus": {}}
                        st.success(f"Vendor {new_v_name} created with {new_v_spoils}% default spoils.")
                    else:
                        st.session_state["vendor_registry"][new_v_id]["name"] = new_v_name
                        st.session_state["vendor_registry"][new_v_id]["default_spoils"] = new_v_spoils / 100.0
                        st.success(f"Vendor {new_v_name} updated successfully.")
                    st.rerun()
                else:
                    st.warning("Please fill out Vendor Name and ID.")

    with col_crm_2:
        st.markdown("### 📦 Add/Update SKU Rates")
        vendor_options = {v_info["name"]: v_id for v_id, v_info in st.session_state["vendor_registry"].items()}
        
        if not vendor_options:
            st.info("Please add a vendor first.")
        else:
            with st.form("add_sku_form", clear_on_submit=True):
                selected_v_name = st.selectbox("Select Vendor", list(vendor_options.keys()))
                new_sku = st.text_input("Costco Item # (SKU)")
                
                col_sku_1, col_sku_2 = st.columns(2)
                with col_sku_1:
                    new_rate = st.number_input("Commission %", min_value=0.0, max_value=100.0, step=0.1)
                with col_sku_2:
                    new_sku_spoils = st.number_input("Spoils Override % (0 = Use Vendor Default)", min_value=0.0, max_value=100.0, step=0.1)
                    
                submit_sku = st.form_submit_button("Save SKU to Matrix")
                
                if submit_sku:
                    if new_sku:
                        v_id = vendor_options[selected_v_name]
                        decimal_rate = new_rate / 100.0
                        decimal_spoils = (new_sku_spoils / 100.0) if new_sku_spoils > 0 else None
                        
                        st.session_state["vendor_registry"][v_id]["skus"][new_sku] = {
                            "comm": decimal_rate,
                            "spoils": decimal_spoils
                        }
                        st.success(f"SKU {new_sku} saved for {selected_v_name}.")
                        st.rerun()
                    else:
                        st.warning("Please enter a valid SKU.")

    st.markdown("---")
    st.markdown("### 🗄️ Current Active Directory")
    
    for v_id, v_info in st.session_state["vendor_registry"].items():
        v_spoils_display = v_info.get('default_spoils', 0.0075) * 100
        with st.expander(f"🏢 {v_info['name']} (ID: {v_id}) | Base Spoils: {v_spoils_display:.2f}%", expanded=True):
            if not v_info["skus"]:
                st.write("No SKUs mapped yet.")
            else:
                formatted_skus = []
                for sku, data in v_info["skus"].items():
                    # Handle data cleanly
                    comm = data['comm'] if isinstance(data, dict) else data
                    spoils = data.get('spoils') if isinstance(data, dict) else None
                    
                    formatted_skus.append({
                        "Costco Item #": sku, 
                        "Commission Rate": f"{comm*100:.1f}%",
                        "Spoils Rate": f"{spoils*100:.2f}% (Override)" if spoils is not None else f"Vendor Default ({v_spoils_display:.2f}%)"
                    })
                    
                st.dataframe(pd.DataFrame(formatted_skus), use_container_width=True, hide_index=True)
